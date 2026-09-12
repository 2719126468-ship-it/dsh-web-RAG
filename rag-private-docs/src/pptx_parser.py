"""PPTX 解析器：提取每页幻灯片里的文字，按页切片。"""
from pathlib import Path
from typing import List

from langchain_core.documents import Document


def parse_pptx(path: Path) -> List[Document]:
    try:
        from pptx import Presentation
    except ImportError:
        print(f"[warn] 未安装 python-pptx，跳过 {path.name}。安装：pip install python-pptx")
        return []

    try:
        prs = Presentation(str(path))
    except Exception as e:
        print(f"[warn] 打开 pptx 失败 {path}: {e}")
        return []

    docs: List[Document] = []
    for idx, slide in enumerate(prs.slides, start=1):
        parts: List[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        parts.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if parts:
            docs.append(Document(
                page_content="\n".join(parts),
                metadata={"source": str(path), "slide": idx},
            ))
    return docs
