"""Excel / CSV 解析器：按行转成文本块。

xlsx 用 openpyxl（纯 Python），csv 用标准库。
"""
import csv
from pathlib import Path
from typing import List

from langchain_core.documents import Document


def parse_xlsx(path: Path) -> List[Document]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        print(f"[warn] 未安装 openpyxl，跳过 {path.name}。安装：pip install openpyxl")
        return []

    try:
        wb = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as e:
        print(f"[warn] 打开 xlsx 失败 {path}: {e}")
        return []

    docs: List[Document] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(c) if c is not None else "" for c in rows[0]]
        lines = [" | ".join(header)]
        for row in rows[1:]:
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                lines.append(" | ".join(cells))
        content = "\n".join(lines)
        if content.strip():
            docs.append(Document(
                page_content=content,
                metadata={"source": str(path), "sheet": sheet_name},
            ))
    wb.close()
    return docs


def parse_csv_file(path: Path) -> List[Document]:
    docs: List[Document] = []
    rows = None
    for enc in ("utf-8-sig", "gbk", "utf-8"):
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                rows = list(csv.reader(f))
            break
        except (UnicodeDecodeError, LookupError):
            continue
        except Exception as e:
            print(f"[warn] 读取 csv 失败 {path}: {e}")
            return []

    if not rows:
        return []
    header = rows[0]
    lines = [" | ".join(header)]
    for row in rows[1:]:
        if any(c.strip() for c in row):
            lines.append(" | ".join(row))
    content = "\n".join(lines)
    if content.strip():
        docs.append(Document(page_content=content, metadata={"source": str(path)}))
    return docs
